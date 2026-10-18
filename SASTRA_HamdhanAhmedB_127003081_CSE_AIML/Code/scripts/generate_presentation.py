import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# -----------------------------------------------------------------------------
# Color Palette Constants
# -----------------------------------------------------------------------------
DARK_NAVY        = RGBColor(13, 25, 45)      # #0D192D - Deep Navy
SECONDARY_NAVY   = RGBColor(20, 38, 64)      # #142640 - Primary Container fill on dark
CARD_NAVY        = RGBColor(16, 30, 52)      # #101E34 - Nested Card fill
NAVY_BORDER      = RGBColor(46, 80, 125)     # #2E507D - Structural Border
ELECTRIC_BLUE    = RGBColor(46, 126, 220)    # #2E7EDC - Primary Accent
CYAN             = RGBColor(52, 198, 210)    # #34C6D2 - Tech Accent Cyan
LIGHT_CYAN       = RGBColor(180, 240, 245)   # #B4F0F5

BG_LIGHT         = RGBColor(248, 250, 252)   # #F8FAFC - Main content background
CARD_LIGHT       = RGBColor(255, 255, 255)   # #FFFFFF - Crisp white cards
BORDER_LIGHT     = RGBColor(226, 232, 240)   # #E2E8F0 - Subtle card border
BORDER_ACCENT    = RGBColor(190, 215, 245)   # #BED7F5

TEXT_DARK        = RGBColor(15, 23, 42)      # #0F172A - High-contrast charcoal text
TEXT_MUTED       = RGBColor(71, 85, 105)     # #475569 - Secondary body text
TEXT_LIGHT       = RGBColor(241, 245, 249)   # #F1F5F9 - High-contrast white on dark
TEXT_LIGHT_MUTED = RGBColor(148, 163, 184)   # #94A3B8 - Secondary text on dark

SUCCESS_GREEN    = RGBColor(16, 185, 129)    # #10B981 - Pass / Validated
WARNING_AMBER    = RGBColor(245, 158, 11)    # #F59E0B - In progress / Revision
DANGER_RED       = RGBColor(225, 29, 72)     # #E11D48 - Problem statement
ACCENT_PURPLE    = RGBColor(139, 92, 246)    # #8B5CF6

FONT_FAMILY = "Segoe UI"

GITHUB_URL   = "https://github.com/hamdhan05/AUTOSAR-AI-Implementation"
STUDENT_NAME = "Hamdhan Ahmed B"
STUDENT_REG  = "127003081"
PROJECT_TITLE = "AUTOSAR AI – High-Level Design Document Intelligence Platform"

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # Reusable Layout & Element Helpers
    # =========================================================================
    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, badge_text, title_text, subtitle_text, is_dark=False):
        # Category Badge
        badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.28))
        tf = badge_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.text = badge_text.upper()
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = CYAN if is_dark else ELECTRIC_BLUE

        # Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.7), Inches(0.55))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.alignment = PP_ALIGN.LEFT
        p_t.text = title_text
        p_t.font.name = FONT_FAMILY
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_LIGHT if is_dark else TEXT_DARK

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.24), Inches(11.7), Inches(0.35))
        tf_s = sub_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        p_s.alignment = PP_ALIGN.LEFT
        p_s.text = subtitle_text
        p_s.font.name = FONT_FAMILY
        p_s.font.size = Pt(11)
        p_s.font.color.rgb = TEXT_LIGHT_MUTED if is_dark else TEXT_MUTED

    def add_footer(slide, current_slide, total_slides=8, is_dark=False):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.08), Inches(10.0), Inches(0.3))
        tf = footer_box.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.text = f"{PROJECT_TITLE}  |  {STUDENT_NAME} (Reg: {STUDENT_REG})"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(8.8)
        p.font.color.rgb = RGBColor(90, 110, 140) if is_dark else RGBColor(140, 155, 175)

        # Right aligned slide number
        page_box = slide.shapes.add_textbox(Inches(10.5), Inches(7.08), Inches(2.033), Inches(0.3))
        tf_p = page_box.text_frame
        tf_p.margin_left = tf_p.margin_top = tf_p.margin_right = tf_p.margin_bottom = 0
        p_p = tf_p.paragraphs[0]
        p_p.alignment = PP_ALIGN.RIGHT
        p_p.text = f"Slide {current_slide} of {total_slides}"
        p_p.font.name = FONT_FAMILY
        p_p.font.size = Pt(9)
        p_p.font.bold = True
        p_p.font.color.rgb = CYAN if is_dark else ELECTRIC_BLUE

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Navy Theme)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, DARK_NAVY)

    # Accent decorative bars
    top_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.55), Inches(0.6), Inches(0.08))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = CYAN
    top_bar.line.fill.background()

    bar2 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.45), Inches(0.55), Inches(1.5), Inches(0.08))
    bar2.fill.solid()
    bar2.fill.fore_color.rgb = ELECTRIC_BLUE
    bar2.line.fill.background()

    # Category Badge
    b1 = s1.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.35))
    tf1 = b1.text_frame
    tf1.margin_left = tf1.margin_top = 0
    p1 = tf1.paragraphs[0]
    p1.alignment = PP_ALIGN.LEFT
    p1.text = "ACADEMIC TECHNICAL PRESENTATION  •  AUTOMOTIVE SOFTWARE & AI SYSTEMS"
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = CYAN

    # Project Title
    t1 = s1.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.5), Inches(1.1))
    tf_t1 = t1.text_frame
    tf_t1.margin_left = tf_t1.margin_top = 0
    p_t1 = tf_t1.paragraphs[0]
    p_t1.alignment = PP_ALIGN.LEFT
    p_t1.text = "AUTOSAR AI"
    p_t1.font.name = FONT_FAMILY
    p_t1.font.size = Pt(44)
    p_t1.font.bold = True
    p_t1.font.color.rgb = RGBColor(255, 255, 255)

    # Subtitle
    sub1 = s1.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(0.6))
    tf_sub1 = sub1.text_frame
    tf_sub1.margin_left = tf_sub1.margin_top = 0
    p_sub1 = tf_sub1.paragraphs[0]
    p_sub1.alignment = PP_ALIGN.LEFT
    p_sub1.text = "High-Level Design Document Intelligence Platform"
    p_sub1.font.name = FONT_FAMILY
    p_sub1.font.size = Pt(22)
    p_sub1.font.color.rgb = LIGHT_CYAN

    # Tagline
    tag1 = s1.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(11.5), Inches(0.45))
    tf_tag1 = tag1.text_frame
    tf_tag1.margin_left = tf_tag1.margin_top = 0
    p_tag1 = tf_tag1.paragraphs[0]
    p_tag1.alignment = PP_ALIGN.LEFT
    p_tag1.text = "“Local-First AI for Evidence-Grounded Automotive Architecture Analysis”"
    p_tag1.font.name = FONT_FAMILY
    p_tag1.font.size = Pt(13)
    p_tag1.font.italic = True
    p_tag1.font.color.rgb = TEXT_LIGHT_MUTED

    # Left Card: Student & Project Details
    card_s = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.55), Inches(6.8), Inches(2.95))
    card_s.fill.solid()
    card_s.fill.fore_color.rgb = SECONDARY_NAVY
    card_s.line.color.rgb = NAVY_BORDER
    card_s.line.width = Pt(1.5)

    card_s_tf = card_s.text_frame
    card_s_tf.vertical_anchor = MSO_ANCHOR.TOP
    card_s_tf.margin_left = Inches(0.35)
    card_s_tf.margin_top = Inches(0.28)
    card_s_tf.margin_right = Inches(0.35)
    card_s_tf.word_wrap = True

    p = card_s_tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.text = "STUDENT & PROJECT SPECIFICATIONS"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = CYAN

    details = [
        ("Candidate Name:", STUDENT_NAME),
        ("Registration Number:", STUDENT_REG),
        ("Domain:", "Automotive Software Engineering & AI Architectures"),
        ("System Paradigm:", "Local-First LLM (Ollama llama3.2:3b) + Persistent Vector Store"),
        ("Verification Standard:", "Strict Source Grounding & Verifiable Citations"),
    ]
    for lbl, val in details:
        p = card_s_tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(6)
        r1 = p.add_run()
        r1.text = f"{lbl} "
        r1.font.name = FONT_FAMILY
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = TEXT_LIGHT_MUTED
        r2 = p.add_run()
        r2.text = val
        r2.font.name = FONT_FAMILY
        r2.font.size = Pt(11)
        r2.font.bold = True
        r2.font.color.rgb = TEXT_LIGHT

    # Right Top Card: Open-Source Repository (Clickable Shape Action)
    card_git = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(3.55), Inches(4.733), Inches(1.35))
    card_git.fill.solid()
    card_git.fill.fore_color.rgb = SECONDARY_NAVY
    card_git.line.color.rgb = CYAN
    card_git.line.width = Pt(1.5)
    card_git.click_action.hyperlink.address = GITHUB_URL

    git_tf = card_git.text_frame
    git_tf.vertical_anchor = MSO_ANCHOR.TOP
    git_tf.margin_left = Inches(0.3)
    git_tf.margin_top = Inches(0.18)
    git_tf.margin_right = Inches(0.3)
    git_tf.word_wrap = True

    p = git_tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.text = "PROJECT REPOSITORY & CODEBASE (CLICK TO OPEN)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = CYAN

    p = git_tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    p.space_before = Pt(4)
    run = p.add_run()
    run.text = GITHUB_URL
    run.font.name = FONT_FAMILY
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = LIGHT_CYAN

    p = git_tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    p.space_before = Pt(3)
    p.text = "Full Stack Implementation • Automated Unit & Integration Tests • Docker Ready"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(8.5)
    p.font.color.rgb = TEXT_LIGHT_MUTED

    # Right Bottom Cards: Key Highlights
    key_badges = [
        ("OFFLINE-FIRST AIR-GAPPED PRIVACY",
         "Ollama (llama3.2:3b) local inference; zero automotive IP leaves local premises.", ELECTRIC_BLUE),
        ("VERIFIED TEST SUITE (100% PASS RATE)",
         "28 / 28 automated tests passing across 12 test modules on Python 3.11.", SUCCESS_GREEN),
    ]
    cur_top = 5.05
    for b_title, b_desc, b_col in key_badges:
        b_shape = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(cur_top), Inches(4.733), Inches(0.7))
        b_shape.fill.solid()
        b_shape.fill.fore_color.rgb = CARD_NAVY
        b_shape.line.color.rgb = NAVY_BORDER
        b_shape.line.width = Pt(1)

        b_tf = b_shape.text_frame
        b_tf.vertical_anchor = MSO_ANCHOR.TOP
        b_tf.margin_left = Inches(0.25)
        b_tf.margin_top = Inches(0.1)
        b_tf.margin_right = Inches(0.25)
        b_tf.word_wrap = True

        p = b_tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.text = b_title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = b_col

        p2 = b_tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        p2.text = b_desc
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = TEXT_LIGHT_MUTED
        cur_top += 0.76

    add_footer(s1, 1, 8, is_dark=True)

    # =========================================================================
    # SLIDE 2: Problem Statement and Project Objectives (Light Theme)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, BG_LIGHT)
    add_header(s2, "Engineering Challenge & Strategic Goals", "Problem Statement and Project Objectives", "Overcoming verification bottlenecks in massive automotive AUTOSAR High-Level Design specifications")

    # Column 1 Header: The Engineering Problem
    p_header = s2.shapes.add_textbox(Inches(0.8), Inches(1.68), Inches(5.65), Inches(0.3))
    p_header.text_frame.margin_left = p_header.text_frame.margin_top = 0
    p = p_header.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.text = "THE ENGINEERING PROBLEM IN AUTOMOTIVE HLD"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = DANGER_RED

    problems = [
        ("Massive Specification Volume & Slow Navigation",
         "AUTOSAR High-Level Design (HLD) specifications span hundreds of dense pages containing deeply nested architectural models. Manual information retrieval requires substantial engineering hours and delays review cycles."),
        ("Fragmented Architectural Traceability",
         "Identifying Software Components (SWCs), Runnables, Ports, and Interface mappings across disparate sections is error-prone. Critical dependencies and signal types are easily overlooked during manual inspection."),
        ("Revision Comparison Drift & Hallucination Risk",
         "Comparing specification revisions line-by-line is tedious and risks silent architecture drift. Meanwhile, generic public cloud LLMs hallucinate non-existent interfaces and lack document grounding.")
    ]

    top_y = 2.05
    for title, desc in problems:
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_y), Inches(5.65), Inches(1.3))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_LIGHT
        c.line.color.rgb = BORDER_LIGHT
        c.line.width = Pt(1)

        # Indicator bar on left of card
        strip = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(top_y + 0.15), Inches(0.08), Inches(1.0))
        strip.fill.solid()
        strip.fill.fore_color.rgb = DANGER_RED
        strip.line.fill.background()

        tf = c.text_frame
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.28)
        tf.margin_top = Inches(0.14)
        tf.margin_right = Inches(0.2)
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.text = title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        p2.space_before = Pt(4)
        p2.text = desc
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_MUTED
        top_y += 1.42

    # Column 2 Header: Project Objectives
    o_header = s2.shapes.add_textbox(Inches(6.75), Inches(1.68), Inches(5.78), Inches(0.3))
    o_header.text_frame.margin_left = o_header.text_frame.margin_top = 0
    p = o_header.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.text = "CORE PROJECT OBJECTIVES & SOLUTIONS"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = SUCCESS_GREEN

    objectives = [
        ("1. Automated Semantic HLD Ingestion",
         "Ingest multi-page PDF specifications, extract text structure page-by-page, and build persistent searchable vector indexes with strict document & page metadata."),
        ("2. Evidence-Grounded Natural-Language Q&A",
         "Allow automotive architects to query complex component behaviors in natural English, retrieving top-k relevant excerpts via local RAG with zero data leakage."),
        ("3. Strict Traceability & Abstention Guardrails",
         "Provide transparent document, section, and page citations for every assertion, with strict LLM abstention directives when source evidence is insufficient."),
        ("4. Architectural Validation & Revision Diffing",
         "Automatically extract SWCs, Runnables, Ports, and Signals; run enterprise rule checks (circular dependencies, insecure IPC); compute baseline-vs-revision deltas.")
    ]

    top_y = 2.05
    for title, desc in objectives:
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(top_y), Inches(5.78), Inches(0.96))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_LIGHT
        c.line.color.rgb = BORDER_LIGHT
        c.line.width = Pt(1)

        # Indicator bar
        strip = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.75), Inches(top_y + 0.12), Inches(0.08), Inches(0.72))
        strip.fill.solid()
        strip.fill.fore_color.rgb = ELECTRIC_BLUE
        strip.line.fill.background()

        tf = c.text_frame
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.28)
        tf.margin_top = Inches(0.1)
        tf.margin_right = Inches(0.2)
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.text = title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        p2.space_before = Pt(2)
        p2.text = desc
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(9.0)
        p2.font.color.rgb = TEXT_MUTED
        top_y += 1.05

    # Bottom Guiding Principle Banner
    prin_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.35), Inches(11.733), Inches(0.58))
    prin_box.fill.solid()
    prin_box.fill.fore_color.rgb = RGBColor(238, 246, 255)
    prin_box.line.color.rgb = RGBColor(190, 220, 250)
    prin_box.line.width = Pt(1)

    tf = prin_box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.12)
    tf.margin_right = Inches(0.3)
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "CORE ENGINEERING PRINCIPLE: "
    run.font.name = FONT_FAMILY
    run.font.size = Pt(9.5)
    run.font.bold = True
    run.font.color.rgb = ELECTRIC_BLUE

    run2 = p.add_run()
    run2.text = "AUTOSAR AI acts as an authoritative analysis assistant and verification co-pilot—accelerating specification discovery while strictly preserving human engineer oversight, review sign-off, and final architectural approval."
    run2.font.name = FONT_FAMILY
    run2.font.size = Pt(9.5)
    run2.font.color.rgb = TEXT_DARK

    add_footer(s2, 2, 8, is_dark=False)

    # =========================================================================
    # SLIDE 3: System Architecture (Technical Navy / High-Contrast Theme)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, DARK_NAVY)
    add_header(s3, "Three-Tier Multi-Service Topology", "System Architecture & Communications Flow", "Decoupled architecture separating presentation, orchestration, persistence, and local AI inference", is_dark=True)

    # Layer 1: Presentation Layer Container
    l1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.78), Inches(3.6), Inches(4.4))
    l1.fill.solid()
    l1.fill.fore_color.rgb = SECONDARY_NAVY
    l1.line.color.rgb = ELECTRIC_BLUE
    l1.line.width = Pt(1.5)

    tf1 = l1.text_frame
    tf1.vertical_anchor = MSO_ANCHOR.TOP
    tf1.margin_left = Inches(0.25)
    tf1.margin_top = Inches(0.2)
    tf1.margin_right = Inches(0.25)
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.text = "PRESENTATION LAYER"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = CYAN

    p_sub = tf1.add_paragraph()
    p_sub.alignment = PP_ALIGN.LEFT
    p_sub.text = "Streamlit Multipage Web UI  •  localhost:8501"
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(9)
    p_sub.font.bold = True
    p_sub.font.color.rgb = LIGHT_CYAN

    ui_modules = [
        ("Dashboard (Page 1)", "Document & entity metrics, subsystem diagnostics"),
        ("Documents (Page 2)", "PDF upload, SHA-256 deduplication, lifecycle tracking"),
        ("Architecture Explorer (Page 3)", "SWC, Runnable, Port, Signal tables & reviews"),
        ("AI Assistant (Page 4)", "RAG conversation, citation inspector, verification"),
        ("Comparisons (Page 5)", "Baseline vs revision diffs, rename detection"),
        ("Findings (Page 6)", "Circular dependencies, insecure IPC resolution"),
        ("Reports & Settings (Page 7)", "PDF/CSV/JSON exports, diagnostic health checks")
    ]
    for m_name, m_desc in ui_modules:
        p = tf1.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(5)
        run = p.add_run()
        run.text = f"• {m_name}: "
        run.font.name = FONT_FAMILY
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = TEXT_LIGHT
        run2 = p.add_run()
        run2.text = m_desc
        run2.font.name = FONT_FAMILY
        run2.font.size = Pt(8.5)
        run2.font.color.rgb = TEXT_LIGHT_MUTED

    # Arrow 1: Streamlit -> FastAPI
    arr1 = s3.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.5), Inches(3.7), Inches(0.4), Inches(0.25))
    arr1.fill.solid()
    arr1.fill.fore_color.rgb = CYAN
    arr1.line.fill.background()

    comm1 = s3.shapes.add_textbox(Inches(4.35), Inches(4.02), Inches(0.7), Inches(0.4))
    comm1.text_frame.margin_left = comm1.text_frame.margin_top = 0
    p = comm1.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "HTTP\nREST"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = CYAN

    # Layer 2: Backend Orchestration Layer Container
    l2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.98), Inches(1.78), Inches(3.9), Inches(4.4))
    l2.fill.solid()
    l2.fill.fore_color.rgb = SECONDARY_NAVY
    l2.line.color.rgb = ELECTRIC_BLUE
    l2.line.width = Pt(1.5)

    tf2 = l2.text_frame
    tf2.vertical_anchor = MSO_ANCHOR.TOP
    tf2.margin_left = Inches(0.25)
    tf2.margin_top = Inches(0.2)
    tf2.margin_right = Inches(0.25)
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.text = "BACKEND ORCHESTRATION"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = CYAN

    p_sub = tf2.add_paragraph()
    p_sub.alignment = PP_ALIGN.LEFT
    p_sub.text = "FastAPI Asynchronous REST API  •  localhost:8000"
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(9)
    p_sub.font.bold = True
    p_sub.font.color.rgb = LIGHT_CYAN

    api_services = [
        ("Document Ingestion Service", "PyMuPDF parsing, bookmark extraction, chunking"),
        ("Hybrid Retrieval Engine", "ChromaDB dense vectors + BM25 keyword + RRF"),
        ("LLM Orchestrator Service", "Context assembly, prompt formatting, abstention"),
        ("Entity & Rule Validator", "Extracts SWC/Ports/Signals; circular dependency checks"),
        ("Grounding & Claim Service", "Claim segmentation & independent evidence checks"),
        ("Comparison & Diff Engine", "Computes structural deltas & rename detection"),
        ("Report Generation Service", "ReportLab PDF generation & CSV sanitization")
    ]
    for s_name, s_desc in api_services:
        p = tf2.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(5)
        run = p.add_run()
        run.text = f"• {s_name}: "
        run.font.name = FONT_FAMILY
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = TEXT_LIGHT
        run2 = p.add_run()
        run2.text = s_desc
        run2.font.name = FONT_FAMILY
        run2.font.size = Pt(8.5)
        run2.font.color.rgb = TEXT_LIGHT_MUTED

    # Arrow 2: FastAPI -> Data & AI Layer
    arr2 = s3.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.98), Inches(3.7), Inches(0.4), Inches(0.25))
    arr2.fill.solid()
    arr2.fill.fore_color.rgb = CYAN
    arr2.line.fill.background()

    comm2 = s3.shapes.add_textbox(Inches(8.83), Inches(4.02), Inches(0.7), Inches(0.4))
    comm2.text_frame.margin_left = comm2.text_frame.margin_top = 0
    p = comm2.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "SQL / IPC\nREST"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = CYAN

    # Layer 3: Data & Local AI Inference Layer Container
    l3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.48), Inches(1.78), Inches(3.05), Inches(4.4))
    l3.fill.solid()
    l3.fill.fore_color.rgb = SECONDARY_NAVY
    l3.line.color.rgb = ELECTRIC_BLUE
    l3.line.width = Pt(1.5)

    tf3 = l3.text_frame
    tf3.vertical_anchor = MSO_ANCHOR.TOP
    tf3.margin_left = Inches(0.22)
    tf3.margin_top = Inches(0.2)
    tf3.margin_right = Inches(0.22)
    tf3.word_wrap = True

    p = tf3.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.text = "DATA & LOCAL AI LAYER"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = CYAN

    p_sub = tf3.add_paragraph()
    p_sub.alignment = PP_ALIGN.LEFT
    p_sub.text = "Zero Cloud Dependency  •  100% Offline"
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(9)
    p_sub.font.bold = True
    p_sub.font.color.rgb = LIGHT_CYAN

    data_comps = [
        ("Relational Database", "SQLite (autosar_ai.db)", "SQLAlchemy ORM + Alembic migrations"),
        ("Vector Store", "ChromaDB (Persistent)", "Persistent HNSW embeddings in ./data/vector_store"),
        ("Embedding Model", "all-MiniLM-L6-v2", "SentenceTransformers (384-dimensional dense vectors)"),
        ("Local Inference LLM", "Ollama (llama3.2:3b)", "Local server at localhost:11434; zero telemetry"),
        ("Storage System", "File System Storage", "Raw PDFs & generated exports in ./data/uploads & exports")
    ]
    for c_type, c_name, c_detail in data_comps:
        p = tf3.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(6)
        run = p.add_run()
        run.text = f"• {c_name}\n"
        run.font.name = FONT_FAMILY
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = TEXT_LIGHT
        run2 = p.add_run()
        run2.text = f"  {c_detail}"
        run2.font.name = FONT_FAMILY
        run2.font.size = Pt(8.0)
        run2.font.color.rgb = TEXT_LIGHT_MUTED

    # Bottom Architectural Highlights Banner
    banner = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.32), Inches(11.733), Inches(0.6))
    banner.fill.solid()
    banner.fill.fore_color.rgb = CARD_NAVY
    banner.line.color.rgb = NAVY_BORDER
    banner.line.width = Pt(1)

    tf = banner.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.12)
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.text = "ARCHITECTURE HIGHLIGHTS: "
    p.font.name = FONT_FAMILY
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = CYAN

    runs = [
        ("Strict Multi-Process Isolation: ", TEXT_LIGHT),
        ("Streamlit (8501) ↔ FastAPI (8000) ↔ Ollama (11434) run as independent processes.  ", TEXT_LIGHT_MUTED),
        ("Air-Gapped Operation: ", TEXT_LIGHT),
        ("Zero external API dependencies; robust privacy for proprietary automotive IP.  ", TEXT_LIGHT_MUTED),
        ("Reproducible: ", TEXT_LIGHT),
        ("Docker & docker-compose enabled.", TEXT_LIGHT_MUTED)
    ]
    for r_txt, r_col in runs:
        r = p.add_run()
        r.text = r_txt
        r.font.name = FONT_FAMILY
        r.font.size = Pt(9.0)
        r.font.color.rgb = r_col

    add_footer(s3, 3, 8, is_dark=True)

    # =========================================================================
    # SLIDE 4: Document Ingestion and Indexing Workflow (Light Theme)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, BG_LIGHT)
    add_header(s4, "Data Processing Pipeline", "Document Ingestion and Indexing Workflow", "Structured multi-stage transformation from raw automotive PDF specifications to searchable knowledge")

    # 6-Step Workflow Cards (Horizontal layout)
    step_width = 1.76
    gap = 0.23
    start_x = 0.8
    y_pos = 1.82
    card_h = 4.22

    steps = [
        ("STEP 1", "Upload & Validate",
         "PDF Upload",
         [
             "Engineer uploads HLD PDF via Streamlit UI",
             "File extension & 50MB size limit validation",
             "SHA-256 fingerprint computed for duplicate check",
             "Document record set to PENDING in SQLite"
         ]),
        ("STEP 2", "Parse & Extract",
         "PyMuPDF Engine",
         [
             "PyMuPDF (fitz) extracts text page-by-page",
             "Physical page numbers preserved for citations",
             "TOC & bookmark hierarchy extracted",
             "Heading context associated with passages"
         ]),
        ("STEP 3", "Chunk & Window",
         "Sliding Window",
         [
             "Sliding window text partitioning",
             "Chunk size: 1,000 characters",
             "Chunk overlap: 150 characters",
             "Preserves sentence and section semantics",
             "Generates unique chunk hashes"
         ]),
        ("STEP 4", "Embeddings",
         "Dense Vectors",
         [
             "SentenceTransformers all-MiniLM-L6-v2",
             "Generates 384-dimensional dense vectors",
             "Attaches document_id, version_label, page_start, page_end",
             "Runs locally on CPU/GPU"
         ]),
        ("STEP 5", "Dual Storage",
         "Persistent Stores",
         [
             "Vectors & metadata saved in ChromaDB",
             "Structured chunk records in SQLite database",
             "Relational foreign keys linked to Document",
             "Document status updated to READY"
         ]),
        ("STEP 6", "Extraction",
         "Entity Discovery",
         [
             "Regex heuristics extract:",
             "  • Software Components (SWC)",
             "  • Runnables (Runnable_*)",
             "  • Interfaces (*_IF)",
             "  • Ports (PPort / RPort)",
             "  • Signals (Signal_*, types)"
         ])
    ]

    for i, (s_num, s_title, s_sub, s_items) in enumerate(steps):
        x = start_x + i * (step_width + gap)

        c = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y_pos), Inches(step_width), Inches(card_h))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_LIGHT
        c.line.color.rgb = BORDER_LIGHT
        c.line.width = Pt(1)

        # Top accent bar on each step
        accent = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y_pos), Inches(step_width), Inches(0.08))
        accent.fill.solid()
        accent.fill.fore_color.rgb = ELECTRIC_BLUE if i < 5 else CYAN
        accent.line.fill.background()

        tf = c.text_frame
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.12)
        tf.margin_top = Inches(0.14)
        tf.margin_right = Inches(0.12)
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.text = s_num
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = ELECTRIC_BLUE

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        p2.text = s_title
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10.5)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_DARK

        p3 = tf.add_paragraph()
        p3.alignment = PP_ALIGN.LEFT
        p3.text = s_sub
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(8.5)
        p3.font.bold = True
        p3.font.color.rgb = CYAN

        for item in s_items:
            p_it = tf.add_paragraph()
            p_it.alignment = PP_ALIGN.LEFT
            p_it.space_before = Pt(4)
            p_it.text = f"• {item}" if not item.startswith("  •") else item
            p_it.font.name = FONT_FAMILY
            p_it.font.size = Pt(8.3)
            p_it.font.color.rgb = TEXT_MUTED

        # Connector arrow between steps
        if i < 5:
            arr = s4.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + step_width + 0.04), Inches(y_pos + 1.8), Inches(0.15), Inches(0.18))
            arr.fill.solid()
            arr.fill.fore_color.rgb = ELECTRIC_BLUE
            arr.line.fill.background()

    # Bottom Clarification & Scope Note
    note_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.68))
    note_box.fill.solid()
    note_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
    note_box.line.color.rgb = BORDER_LIGHT
    note_box.line.width = Pt(1)

    tf_n = note_box.text_frame
    tf_n.vertical_anchor = MSO_ANCHOR.TOP
    tf_n.margin_left = Inches(0.25)
    tf_n.margin_top = Inches(0.12)
    tf_n.margin_right = Inches(0.25)
    tf_n.word_wrap = True

    p = tf_n.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "IMPLEMENTATION STATUS & ARCHITECTURAL SCOPE NOTE:  "
    run.font.name = FONT_FAMILY
    run.font.size = Pt(9.5)
    run.font.bold = True
    run.font.color.rgb = TEXT_DARK

    run2 = p.add_run()
    run2.text = "Current implementation actively extracts Software Components, Runnables, Interfaces, Ports (PPort/RPort), and Signals via pattern-based heuristic extractors with confidence scores (0.88–0.92). Granular AST-based ARXML schema validation, complex bus packaging, and schedule table extraction are designated for Phase 2."
    run2.font.name = FONT_FAMILY
    run2.font.size = Pt(9)
    run2.font.color.rgb = TEXT_MUTED

    add_footer(s4, 4, 8, is_dark=False)

    # =========================================================================
    # SLIDE 5: AI Assistant and RAG Working (Light Theme)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, BG_LIGHT)
    add_header(s5, "Retrieval-Augmented Generation Architecture", "AI Assistant and RAG Working", "Local inference pipeline delivering evidence-grounded answers with verifiable citations and abstention safeguards")

    # 4-Stage Horizontal Pipeline Cards
    pipe_w = 2.75
    pipe_gap = 0.24
    start_x = 0.8
    y_pos = 1.82
    pipe_h = 3.65

    stages = [
        ("STAGE 1: RETRIEVAL", "Hybrid Semantic Search",
         ELECTRIC_BLUE,
         [
             "Engineer enters natural-language engineering query in UI",
             "Query encoded into 384-d vector via all-MiniLM-L6-v2",
             "ChromaDB performs dense cosine vector search (top_k=5)",
             "BM25 keyword search scores candidate document passages",
             "Reciprocal Rank Fusion (RRF) merges dense and keyword ranks",
             "Extracts top hits with document, page, and section metadata"
         ]),
        ("STAGE 2: PROMPTING", "Context Assembly",
         SECONDARY_NAVY,
         [
             "FastAPI builds strict automotive architecture system prompt",
             "Assembles top-k retrieved chunks into structured context window",
             "Enforces explicit citation tagging: [Doc: HLD_v1.0, Page 4]",
             "Strict Abstention Mandate: If evidence is insufficient, LLM must state lack of evidence rather than hallucinating"
         ]),
        ("STAGE 3: INFERENCE", "Local LLM Synthesis",
         CYAN,
         [
             "Dispatches assembled context prompt to local Ollama runtime",
             "Active model: llama3.2:3b (or Qwen2.5:3b) at localhost:11434",
             "Low-temperature deterministic inference (~0.2 temperature)",
             "Zero telemetry or cloud traffic; completely air-gapped",
             "Synthesizes grounded technical answer referencing citations"
         ]),
        ("STAGE 4: VALIDATION", "Grounding & Display",
         SUCCESS_GREEN,
         [
             "Claim Segmentation breaks response into individual assertions",
             "Independent claim verification matches assertions to excerpts",
             "Computes Grounding Score: (Grounded Claims / Total Claims) %",
             "UI renders response with clickable citation source excerpts",
             "Flags 'Review Required' badge if Grounding Score is below 80%"
         ])
    ]

    for i, (st_name, st_title, st_col, st_items) in enumerate(stages):
        x = start_x + i * (pipe_w + pipe_gap)

        c = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y_pos), Inches(pipe_w), Inches(pipe_h))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_LIGHT
        c.line.color.rgb = BORDER_LIGHT
        c.line.width = Pt(1)

        # Header accent bar
        hbar = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y_pos), Inches(pipe_w), Inches(0.08))
        hbar.fill.solid()
        hbar.fill.fore_color.rgb = st_col
        hbar.line.fill.background()

        tf = c.text_frame
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.18)
        tf.margin_top = Inches(0.16)
        tf.margin_right = Inches(0.18)
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.text = st_name
        p.font.name = FONT_FAMILY
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = st_col

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        p2.text = st_title
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_DARK

        for item in st_items:
            p_it = tf.add_paragraph()
            p_it.alignment = PP_ALIGN.LEFT
            p_it.space_before = Pt(4)
            p_it.text = f"• {item}"
            p_it.font.name = FONT_FAMILY
            p_it.font.size = Pt(8.4)
            p_it.font.color.rgb = TEXT_MUTED

        # Arrow to next stage
        if i < 3:
            arr = s5.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + pipe_w + 0.05), Inches(y_pos + 1.6), Inches(0.14), Inches(0.2))
            arr.fill.solid()
            arr.fill.fore_color.rgb = ELECTRIC_BLUE
            arr.line.fill.background()

    # Bottom Safety & Traceability Pillars (3 Cards)
    b_y = 5.65
    b_w = 3.75
    b_gap = 0.24

    pillars = [
        ("Evidence-Grounded Synthesis",
         "Strict System Prompt constraint: LLM answers only using retrieved passages. If asked about components outside the document, it refrains from speculating.",
         ELECTRIC_BLUE),
        ("Granular Traceability Citations",
         "Every assertion is supported by document name, version label, and exact physical page number, allowing engineers to verify claims against the original PDF in seconds.",
         CYAN),
        ("Deterministic Abstention Safeguard",
         "When retrieved context has similarity below threshold (< 0.5) or lacks key entities, the assistant outputs: 'The available documents do not provide sufficient evidence.'",
         SUCCESS_GREEN)
    ]

    for j, (p_title, p_desc, p_col) in enumerate(pillars):
        px = start_x + j * (b_w + b_gap)
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(px), Inches(b_y), Inches(b_w), Inches(1.22))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(248, 250, 252)
        card.line.color.rgb = BORDER_LIGHT
        card.line.width = Pt(1)

        bar = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(px), Inches(b_y + 0.12), Inches(0.06), Inches(0.95))
        bar.fill.solid()
        bar.fill.fore_color.rgb = p_col
        bar.line.fill.background()

        tf = card.text_frame
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.1)
        tf.margin_right = Inches(0.15)
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.text = p_title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        p2.space_before = Pt(3)
        p2.text = p_desc
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = TEXT_MUTED

    add_footer(s5, 5, 8, is_dark=False)

    # =========================================================================
    # SLIDE 6: Implemented Functional Modules (Light Theme)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, BG_LIGHT)
    add_header(s6, "Implemented Application Modules", "Comprehensive Functional Modules", "Six verified functional subsystems spanning document lifecycle, analysis, validation, and enterprise reporting")

    # 6-Card Grid (3 columns x 2 rows)
    col_w = 3.75
    col_gap = 0.24
    row_h = 2.25
    row_gap = 0.22
    start_x = 0.8
    start_y = 1.82

    modules = [
        ("01", "DASHBOARD & HEALTH DIAGNOSTICS",
         "Centralized Metrics & Subsystem Telemetry",
         [
             "Aggregated document counts, total chunk volume, and entity breakdown",
             "Active LLM provider diagnostics: model name, base URL, latency probe",
             "Vector store health indicator and ChromaDB collection status",
             "Audit logs tracking recent ingestion, queries, and review actions"
         ], ELECTRIC_BLUE),
        ("02", "DOCUMENT REPOSITORY & LIFECYCLE",
         "PDF Ingestion, Deduplication & Management",
         [
             "Single/batch PDF upload with real-time file size & type validation",
             "SHA-256 fingerprint deduplication prevents duplicate indexing",
             "Background processing lifecycle: PENDING → PROCESSING → READY",
             "Document metadata inspector (page count, upload timestamp, version)"
         ], CYAN),
        ("03", "ARCHITECTURE EXPLORER",
         "Extracted Components, Ports & Review Actions",
         [
             "Categorized catalog: Software Components, Runnables, Interfaces, Ports, Signals",
             "Confidence scoring per extracted entity (0.88 to 0.92)",
             "Source excerpt & physical page number mapping for auditability",
             "Interactive engineer review actions (Approve, Reject, Flag for Review)"
         ], ACCENT_PURPLE),
        ("04", "AI ASSISTANT (RAG ENGINE)",
         "Natural-Language Q&A with Evidence Citations",
         [
             "Conversational interface grounded strictly in indexed HLD documents",
             "Document-scoped querying allowing filtering by specific HLD version",
             "Clickable citation pills displaying exact source excerpts and pages",
             "Claim verification drawer displaying Grounding Score and confidence"
         ], ELECTRIC_BLUE),
        ("05", "COMPARISONS & REVISION DIFF",
         "Baseline vs Revision Architectural Deltas",
         [
             "Automated structural comparison between document versions (e.g., v1.0 vs v2.0)",
             "Categorizes architectural deltas: Added, Removed, and Modified entities",
             "Signal data type change detection (e.g., uint16 modified to uint32)",
             "Entity rename detection tracking modified component designations"
         ], WARNING_AMBER),
        ("06", "FINDINGS, VALIDATION & REPORTS",
         "Automated Rule Checks & Multi-Format Exports",
         [
             "Rule engine detects circular dependencies, insecure IPC, unbounded arrays",
             "Unresolved interface reference identification and signal type conflicts",
             "One-click ReportLab PDF generation, structured JSON, and sanitized CSV exports",
             "CSV formula injection defense (=, +, -, @ sanitized) preventing spreadsheet exploits"
         ], SUCCESS_GREEN)
    ]

    for idx, (m_id, m_title, m_sub, m_points, m_color) in enumerate(modules):
        col = idx % 3
        row = idx // 3
        x = start_x + col * (col_w + col_gap)
        y = start_y + row * (row_h + row_gap)

        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(col_w), Inches(row_h))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_LIGHT
        card.line.color.rgb = BORDER_LIGHT
        card.line.width = Pt(1)

        # Left color strip
        strip = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.15), Inches(0.06), Inches(row_h - 0.3))
        strip.fill.solid()
        strip.fill.fore_color.rgb = m_color
        strip.line.fill.background()

        tf = card.text_frame
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.12)
        tf.margin_right = Inches(0.15)
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run_id = p.add_run()
        run_id.text = f"{m_id}. "
        run_id.font.name = FONT_FAMILY
        run_id.font.size = Pt(10)
        run_id.font.bold = True
        run_id.font.color.rgb = m_color

        run_t = p.add_run()
        run_t.text = m_title
        run_t.font.name = FONT_FAMILY
        run_t.font.size = Pt(10)
        run_t.font.bold = True
        run_t.font.color.rgb = TEXT_DARK

        p_sub = tf.add_paragraph()
        p_sub.alignment = PP_ALIGN.LEFT
        p_sub.text = m_sub
        p_sub.font.name = FONT_FAMILY
        p_sub.font.size = Pt(8.5)
        p_sub.font.color.rgb = TEXT_MUTED

        for pt in m_points:
            p_pt = tf.add_paragraph()
            p_pt.alignment = PP_ALIGN.LEFT
            p_pt.space_before = Pt(3)
            p_pt.text = f"• {pt}"
            p_pt.font.name = FONT_FAMILY
            p_pt.font.size = Pt(8.2)
            p_pt.font.color.rgb = TEXT_DARK

    # Bottom notice on verified scope
    foot_box = s6.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.4))
    foot_box.text_frame.margin_left = foot_box.text_frame.margin_top = 0
    p = foot_box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.text = "Verified Scope: All 6 modules are operational in the existing codebase with corresponding Streamlit pages and FastAPI endpoints. Advanced ARXML round-tripping and multi-agent reasoning are classified as future scope."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(9)
    p.font.italic = True
    p.font.color.rgb = TEXT_MUTED

    add_footer(s6, 6, 8, is_dark=False)

    # =========================================================================
    # SLIDE 7: Testing, Verification and Benchmark Results (Light Theme)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, BG_LIGHT)
    add_header(s7, "Empirical Validation & Test Results", "Testing, Verification & Quality Assurance", "Empirically verified test results across all functional modules, security controls, and RAG retrieval pipelines")

    # Top Metrics Dashboard (4 Cards)
    m_w = 2.75
    m_gap = 0.24
    m_y = 1.8
    metrics = [
        ("28 / 28", "Automated Tests Passed", "100% Pass Rate across all 12 test suites", SUCCESS_GREEN),
        ("0", "Test Failures / Errors", "Zero regressions detected in verification audit", SUCCESS_GREEN),
        ("35", "Evaluation Queries", "Curated benchmark dataset for retrieval evaluation", ELECTRIC_BLUE),
        ("100%", "Offline Capable", "Verified local execution with zero internet connectivity", CYAN),
    ]

    for i, (m_val, m_lbl, m_sub, m_col) in enumerate(metrics):
        x = start_x + i * (m_w + m_gap)
        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(m_y), Inches(m_w), Inches(1.18))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_LIGHT
        card.line.color.rgb = BORDER_LIGHT
        card.line.width = Pt(1)

        # Top border accent
        top_acc = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(m_y), Inches(m_w), Inches(0.06))
        top_acc.fill.solid()
        top_acc.fill.fore_color.rgb = m_col
        top_acc.line.fill.background()

        tf = card.text_frame
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.12)
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.text = m_val
        p.font.name = FONT_FAMILY
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = m_col

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        p2.text = m_lbl
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_DARK

        p3 = tf.add_paragraph()
        p3.alignment = PP_ALIGN.LEFT
        p3.text = m_sub
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(8.0)
        p3.font.color.rgb = TEXT_MUTED

    # Test Suite Breakdown (4 Functional Areas)
    tb_y = 3.15
    tb_w = 5.75
    tb_h = 1.5

    test_areas = [
        ("Document Ingestion & Chunking Pipeline (5 Tests Passed)",
         [
             "test_chunking_and_entity_extraction: Verifies 1000-char chunking with 150-char overlap",
             "test_full_document_ingestion_and_cleanup: Ingests PDF, tests ChromaDB indexing & deletion",
             "test_create_project, test_get_projects, test_upload_invalid_file: API input validation"
         ], ELECTRIC_BLUE),
        ("Entity, Port/Signal & Rule Validation (6 Tests Passed)",
         [
             "test_ports_and_signals_extraction: Confirms PPort/RPort, direction & typed signals (float32)",
             "test_mermaid_generation_and_dependency_traversal: Verifies graph dependency generation",
             "test_findings_and_entity_review: Validates circular dependency and insecure IPC rule triggers"
         ], ACCENT_PURPLE),
        ("Local Ollama LLM, Grounding & Traceability (7 Tests Passed)",
         [
             "test_ollama_connection_diagnostic & test_active_provider_info: Confirms Ollama health probe",
             "test_ask_question_insufficient_evidence: Verifies deterministic abstention guardrail",
             "test_claim_segmentation, test_evidence_verification, test_grounding_score: Verifies claim checks"
         ], SUCCESS_GREEN),
        ("Revision Comparison, Reports & Security (10 Tests Passed)",
         [
             "test_csv_formula_injection_sanitization: Sanitizes '=cmd', '+', '-', '@' against injection",
             "test_rename_detection_and_data_type_diff: Detects entity renames and signal type deltas",
             "test_report_generation_and_download, test_governance, test_audit: Validates export workflows"
         ], WARNING_AMBER)
    ]

    for idx, (t_title, t_items, t_col) in enumerate(test_areas):
        col = idx % 2
        row = idx // 2
        x = start_x + col * (tb_w + 0.23)
        y = tb_y + row * (tb_h + 0.15)

        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(tb_w), Inches(tb_h))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_LIGHT
        card.line.color.rgb = BORDER_LIGHT
        card.line.width = Pt(1)

        strip = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.12), Inches(0.06), Inches(tb_h - 0.24))
        strip.fill.solid()
        strip.fill.fore_color.rgb = t_col
        strip.line.fill.background()

        tf = card.text_frame
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.12)
        tf.margin_right = Inches(0.15)
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.text = t_title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK

        for item in t_items:
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.LEFT
            p2.space_before = Pt(3)
            p2.text = f"• {item}"
            p2.font.name = FONT_FAMILY
            p2.font.size = Pt(8.3)
            p2.font.color.rgb = TEXT_MUTED

    # End-to-End Walkthrough Box
    e2e_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.32), Inches(11.733), Inches(0.62))
    e2e_box.fill.solid()
    e2e_box.fill.fore_color.rgb = RGBColor(238, 246, 255)
    e2e_box.line.color.rgb = RGBColor(190, 220, 250)
    e2e_box.line.width = Pt(1)

    tf_e = e2e_box.text_frame
    tf_e.vertical_anchor = MSO_ANCHOR.TOP
    tf_e.margin_left = Inches(0.25)
    tf_e.margin_top = Inches(0.1)
    tf_e.word_wrap = True

    p = tf_e.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "VERIFIED END-TO-END WORKFLOW EXAMPLE:  "
    run.font.name = FONT_FAMILY
    run.font.size = Pt(9.5)
    run.font.bold = True
    run.font.color.rgb = ELECTRIC_BLUE

    run2 = p.add_run()
    run2.text = "AUTOSAR HLD PDF Upload (Page 4) → PyMuPDF Extraction → ChromaDB Embedding Indexing → Engineer Query: 'What interfaces does MotorControl_SWC provide?' → Top-k Retrieval → Local Ollama Inference → Verified Cited Answer with Page 4 excerpt and 92% Grounding Score."
    run2.font.name = FONT_FAMILY
    run2.font.size = Pt(8.8)
    run2.font.color.rgb = TEXT_DARK

    add_footer(s7, 7, 8, is_dark=False)

    # =========================================================================
    # SLIDE 8: Limitations, Future Enhancements and Conclusion (Dark Navy Theme)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, DARK_NAVY)
    add_header(s8, "Roadmap & Strategic Conclusion", "Four-Phase Evolution & Strategic Conclusion", "From functional prototype to an enterprise-grade automotive architecture intelligence platform", is_dark=True)

    # 4-Phase Roadmap Cards (Horizontal)
    p_w = 2.75
    p_gap = 0.24
    p_y = 1.82
    p_h = 3.65

    phases = [
        ("PHASE 1", "Traceability & Governance",
         ELECTRIC_BLUE,
         "CURRENT BASELINE & REFINEMENTS\n\n"
         "• Formal document version tagging & governance metadata\n"
         "• Claim-level answer segmentation & independent verification\n"
         "• Inline Markdown, JSON, and sanitized CSV exports\n"
         "• Document-scoped query filtering\n"
         "• Grounding Score & Review Recommendation badge"),
        ("PHASE 2", "Architecture Depth",
         CYAN,
         "TOPOLOGY & SCHEMA EXPANSION\n\n"
         "• Full AST-based Port, Signal & Interface extraction\n"
         "• Interactive architecture topology graphs via Mermaid & Cytoscape\n"
         "• Dependency traversal & change-impact propagation\n"
         "• Bidirectional ARXML schema round-tripping & synchronization\n"
         "• Bus frame packaging & signal timing analysis"),
        ("PHASE 3", "Enterprise Governance",
         ACCENT_PURPLE,
         "COMPLIANCE & TEAM COLLABORATION\n\n"
         "• Formal engineer sign-off workflow for AI-assisted analysis\n"
         "• Immutable audit trails & persistent Q&A review history\n"
         "• Per-component architectural compliance scorecards\n"
         "• Multi-user role-based access control (RBAC)\n"
         "• Integration with JIRA and enterprise PLM tools"),
        ("PHASE 4", "Evaluation & Multi-Agent",
         SUCCESS_GREEN,
         "BENCHMARKING & ADVANCED AGENTS\n\n"
         "• Automated benchmark execution against 35-query dataset\n"
         "• Retrieval metrics tracking (MRR, Hit Rate@1, latency)\n"
         "• Hybrid BM25 + Dense retrieval dynamic weight tuning\n"
         "• Multi-agent architectural critique & consistency validation\n"
         "• ISO 26262 functional safety traceability mapping")
    ]

    for i, (ph_num, ph_title, ph_col, ph_body) in enumerate(phases):
        x = start_x + i * (p_w + p_gap)
        card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(p_y), Inches(p_w), Inches(p_h))
        card.fill.solid()
        card.fill.fore_color.rgb = SECONDARY_NAVY
        card.line.color.rgb = NAVY_BORDER
        card.line.width = Pt(1.5)

        # Top accent bar
        top_b = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(p_y), Inches(p_w), Inches(0.08))
        top_b.fill.solid()
        top_b.fill.fore_color.rgb = ph_col
        top_b.line.fill.background()

        tf = card.text_frame
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.16)
        tf.margin_top = Inches(0.15)
        tf.margin_right = Inches(0.16)
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.text = ph_num
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = ph_col

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        p2.text = ph_title
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_LIGHT

        p3 = tf.add_paragraph()
        p3.alignment = PP_ALIGN.LEFT
        p3.space_before = Pt(8)
        p3.text = ph_body
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(8.5)
        p3.font.color.rgb = TEXT_LIGHT_MUTED

    # Bottom Strategic Conclusion Banner
    c_banner = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.62), Inches(11.733), Inches(0.85))
    c_banner.fill.solid()
    c_banner.fill.fore_color.rgb = CARD_NAVY
    c_banner.line.color.rgb = ELECTRIC_BLUE
    c_banner.line.width = Pt(1.5)

    tf_c = c_banner.text_frame
    tf_c.vertical_anchor = MSO_ANCHOR.TOP
    tf_c.margin_left = Inches(0.25)
    tf_c.margin_top = Inches(0.12)
    tf_c.margin_right = Inches(0.25)
    tf_c.word_wrap = True

    p = tf_c.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "STRATEGIC CONCLUSION:  "
    run.font.name = FONT_FAMILY
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = CYAN

    run2 = p.add_run()
    run2.text = "AUTOSAR AI successfully bridges the gap between massive, static High-Level Design PDF specifications and dynamic engineering workflows. By uniting local RAG retrieval, rule-based architecture validation, automated revision diffing, and strict citation traceability, the platform provides automotive architects with verifiable, evidence-grounded insights while preserving complete offline data sovereignty."
    run2.font.name = FONT_FAMILY
    run2.font.size = Pt(9.0)
    run2.font.color.rgb = TEXT_LIGHT

    # Final Clickable GitHub & Student Banner (Pill Button)
    git_btn = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.2), Inches(6.58), Inches(8.933), Inches(0.42))
    git_btn.fill.solid()
    git_btn.fill.fore_color.rgb = SECONDARY_NAVY
    git_btn.line.color.rgb = CYAN
    git_btn.line.width = Pt(1.2)
    git_btn.click_action.hyperlink.address = GITHUB_URL

    tf_f = git_btn.text_frame
    tf_f.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_f.margin_left = tf_f.margin_top = tf_f.margin_right = tf_f.margin_bottom = 0
    p = tf_f.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER

    r1 = p.add_run()
    r1.text = f"{STUDENT_NAME}   |   Registration No: {STUDENT_REG}   |   GitHub: "
    r1.font.name = FONT_FAMILY
    r1.font.size = Pt(10.5)
    r1.font.bold = True
    r1.font.color.rgb = TEXT_LIGHT_MUTED

    r2 = p.add_run()
    r2.text = GITHUB_URL
    r2.font.name = FONT_FAMILY
    r2.font.size = Pt(10.5)
    r2.font.bold = True
    r2.font.color.rgb = LIGHT_CYAN

    add_footer(s8, 8, 8, is_dark=True)

    return prs

if __name__ == "__main__":
    out_dir = r"h:\tech_pulse_project"
    out_file = os.path.join(out_dir, "AUTOSAR_AI_Hamdhan_Ahmed_B_Presentation.pptx")
    prs = create_presentation()
    prs.save(out_file)
    print(f"Presentation successfully created at: {out_file}")
    
    # Also save a copy inside autosar-ai folder
    copy_file = os.path.join(out_dir, "autosar-ai", "AUTOSAR_AI_Hamdhan_Ahmed_B_Presentation.pptx")
    prs.save(copy_file)
    print(f"Copy saved at: {copy_file}")
